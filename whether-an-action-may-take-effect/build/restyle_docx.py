"""Apply the Loomground palette to a paper .docx, in place.

White ground, black type, violet headings and links, code in violet mono on #F8F8F8,
and no colour outside the palette: pandoc's syntax-highlighting colours and the Word
theme colours are replaced, so nothing falls back to Word's defaults.

Run: python3 build/restyle_docx.py <file.docx> [<file.docx> ...]
"""
import re, sys, zipfile
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_docx import font, rule, INK, VIOLET, META, SANS, DISPLAY, MONO  # noqa: E402

SURFACE = "F8F8F8"
# Word theme slots, set to the palette (black tints for the neutral slots).
THEME = {"dk1": "111111", "lt1": "FFFFFF", "dk2": "111111", "lt2": SURFACE,
         "accent1": "5B5B9A", "accent2": "4A4A80", "accent3": "555555", "accent4": "9A9A9A",
         "accent5": "DDDDDD", "accent6": "E8E8F5", "hlink": "5B5B9A", "folHlink": "4A4A80"}


def shade(style, fill):
    ppr = style.element.get_or_add_pPr()
    for old in ppr.findall(qn("w:shd")):
        ppr.remove(old)
    shd = OxmlElement("w:shd")
    for k, v in (("val", "clear"), ("color", "auto"), ("fill", fill)):
        shd.set(qn(f"w:{k}"), v)
    ppr.append(shd)


def styles(d):
    # look styles up by their fixed ids: display names differ between files ("Heading 1"
    # in one, the built-in lower-case "heading 1" in another)
    s = {st.style_id: st for st in d.styles if getattr(st, "style_id", None)}
    for n in ("Normal", "BodyText", "FirstParagraph", "Compact", "BlockText", "Abstract"):
        if n in s: font(s[n], SANS, 10.5, INK)
    if "Title" in s: font(s["Title"], DISPLAY, 24, INK, True)
    if "AbstractTitle" in s: font(s["AbstractTitle"], DISPLAY, 10.5, VIOLET, True)
    if "Subtitle" in s: font(s["Subtitle"], SANS, 12, META, False)
    for n, size, colour in (("Heading1", 20, INK), ("Heading2", 15, VIOLET), ("Heading3", 12, VIOLET),
                            ("Heading4", 10.5, VIOLET), ("Heading5", 10.5, VIOLET), ("Heading6", 10.5, VIOLET)):
        if n in s: font(s[n], DISPLAY, size, colour, True)
        if n + "Char" in s: font(s[n + "Char"], DISPLAY, None, colour)
    for n in ("Caption", "ImageCaption", "TableCaption"):
        if n in s: font(s[n], MONO, 8.5, META, False)
    if "Hyperlink" in s: font(s["Hyperlink"], SANS, None, VIOLET)
    if "VerbatimChar" in s: font(s["VerbatimChar"], MONO, 8.5, INK)
    if "SourceCode" in s:
        font(s["SourceCode"], MONO, 8.5, VIOLET)
        shade(s["SourceCode"], SURFACE)
    for sid, st in s.items():   # pandoc's highlighting tokens: one colour, the violet of the code block
        if sid.endswith("Tok"):
            font(st, MONO, None, VIOLET)


def detheme(d):
    """a theme reference beats an explicit value in Word: drop every theme font and colour
    reference from the styles, so the explicit Loomground values are the ones that render"""
    root = d.styles.element
    for el in root.iter(qn("w:rFonts")):
        for k in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            el.attrib.pop(qn(k), None)
    palette = {"111111", "FFFFFF", "5B5B9A", "4A4A80", "E8E8F5", "F8F8F8", "F4F4F4", "F0F0F0",
               "DDDDDD", "555555", "9A9A9A", "000000", "AUTO"}
    for el in root.iter(qn("w:color")):
        for k in ("w:themeColor", "w:themeShade", "w:themeTint"):
            el.attrib.pop(qn(k), None)
        val = (el.get(qn("w:val")) or "").upper()
        if val and val not in palette:
            # a left-over theme fallback: a coloured one becomes the violet, a grey the metadata grey
            r, g, b = (int(val[i:i + 2], 16) for i in (0, 2, 4))
            el.set(qn("w:val"), "5B5B9A" if max(r, g, b) - min(r, g, b) > 24 else "555555")


def paragraphs(d):
    for par in d.paragraphs:
        if par.style.name == "Source Code":
            for run in par.runs:
                run.font.name = MONO
                run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), MONO)
    for t in d.tables:
        if t._tbl.tblPr.find(qn("w:tblBorders")) is None:
            rule(t)


def theme(path):
    """rewrite the theme colour slots in place, keeping every other part byte for byte"""
    tmp = path.with_suffix(".restyle.tmp")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith("word/theme/") and item.filename.endswith(".xml"):
                x = data.decode("utf-8")
                for slot, hexv in THEME.items():
                    x = re.sub(rf"(<a:{slot}>\s*)<a:(?:srgbClr|sysClr)[^>]*?/>",
                               rf'\1<a:srgbClr val="{hexv}"/>', x)
                data = x.encode("utf-8")
            zout.writestr(item, data)
    tmp.replace(path)


def restyle(path):
    path = Path(path)
    d = Document(path)
    styles(d)
    detheme(d)
    paragraphs(d)
    d.save(path)
    theme(path)
    return path


if __name__ == "__main__":
    for p in sys.argv[1:]:
        print(restyle(p))

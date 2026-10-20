"""paper.md → paper.docx (pandoc with a Loomground-styled reference document).
Run: python3 build/build_docx.py                        (Paper I)
     python3 build/build_docx.py series/F-5d-nd.md --draft  (any paper; --draft shows PENDING/OPEN/CHECK-REF tags)"""
import re, subprocess
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
BUILD = ROOT / "build"
INK, VIOLET, META = RGBColor(0x11, 0x11, 0x11), RGBColor(0x5B, 0x5B, 0x9A), RGBColor(0x55, 0x55, 0x55)
SANS, DISPLAY, MONO = "Archivo", "Archivo", "IBM Plex Mono"   # the fonts the kit installs


def font(style, name, size=None, colour=None, bold=None):
    f = style.font
    f.name = name
    rpr = style.element.get_or_add_rPr()
    rpr.get_or_add_rFonts().set(qn("w:eastAsia"), name)
    for k in ("w:ascii", "w:hAnsi", "w:cs"):
        rpr.get_or_add_rFonts().set(qn(k), name)
    if size: f.size = Pt(size)
    if colour is not None: f.color.rgb = colour
    if bold is not None: f.bold = bold


def rule(table, header=0):
    """header rule in ink, hairline rows in the rule grey — the page style's table"""
    pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge, size, colour in (("top", 12, "111111"), ("bottom", 8, "111111"), ("insideH", 4, "DDDDDD")):
        e = OxmlElement(f"w:{edge}")
        for k, v in (("val", "single"), ("sz", str(size)), ("space", "0"), ("color", colour)):
            e.set(qn(f"w:{k}"), v)
        borders.append(e)
    pr.append(borders)
    for cell in table.rows[header].cells:
        for par in cell.paragraphs:
            for run in par.runs:
                run.font.bold = True


def reference():
    ref = BUILD / "reference.docx"
    subprocess.run(["pandoc", "-o", str(ref), "--print-default-data-file", "reference.docx"], check=True)
    d = Document(ref)
    s = d.styles
    for n in ("Normal", "Body Text", "First Paragraph", "Compact"):
        if n in s: font(s[n], SANS, 10.5, INK)
    font(s["Title"], DISPLAY, 24, INK, True)
    if "Subtitle" in s: font(s["Subtitle"], SANS, 12, META, False)
    font(s["Heading 1"], DISPLAY, 20, INK, True)
    font(s["Heading 2"], DISPLAY, 15, VIOLET, True)
    font(s["Heading 3"], DISPLAY, 12, INK, True)
    for n in ("Source Code", "Verbatim Char"):
        if n in s: font(s[n], MONO, 8.5, INK)
    for n in ("Caption", "Image Caption", "Table Caption"):
        if n in s: font(s[n], MONO, 8.5, META, False)
    if "Hyperlink" in s: font(s["Hyperlink"], SANS, None, VIOLET)
    d.save(ref)
    return ref


def draft_marks(s):
    """make the anticipation tags visible: a reader of a draft must see which claims wait on the build"""
    s = re.sub(r"<!--PENDING:phase-(\d+)-->", r" ⟦pending: phase \1⟧", s)
    s = re.sub(r"<!--OPEN:?([^>]*?)-->", lambda m: " ⟦open" + (": " + m.group(1).strip() if m.group(1).strip() else "") + "⟧", s)
    return re.sub(r"<!--CHECK-REF-->", " ⟦check ref⟧", s)


def source(md=None, draft=False):
    s = (md or ROOT / "paper.md").read_text()
    if draft:
        s = draft_marks(s)
    s = re.sub(r"<!--.*?-->\n?", "", s)
    s = re.sub(r"```latex\n(.*?)\n```", lambda m: "$$" + m.group(1) + "$$", s, flags=re.S)
    s = re.sub(r"\]\((?:\.\./)?(figures/[^)]+)\.svg\)", r"](\1.png)", s)
    tmp = BUILD / "_paper.md"
    tmp.write_text(s)
    return tmp


def main(argv=()):
    args = [a for a in argv if not a.startswith("--")]
    src = (ROOT / args[0]).resolve() if args else ROOT / "paper.md"
    draft = "--draft" in argv
    title = next((l[2:].strip() for l in src.read_text().splitlines() if l.startswith("# ")), src.stem)
    ref, md = reference(), source(src, draft)
    out = src.with_suffix(".docx") if args else ROOT / "paper.docx"
    subprocess.run(["pandoc", str(md), "-f", "markdown+tex_math_dollars", "-t", "docx", "--syntax-highlighting=none",
                    "--reference-doc", str(ref), "--resource-path", str(ROOT),
                    "--metadata", "title=" + (title.split(":")[0] if args else "Loomground Governance"),
                    "-o", str(out)], check=True, cwd=ROOT)
    md.unlink()
    ref.unlink()
    d = Document(out)
    # the H1 in the source is the real title; pandoc's metadata title only names the file
    first = d.paragraphs[0]
    if first.style.name == "Title":
        first._element.getparent().remove(first._element)
    for sec in d.sections:
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2.2)
        sec.top_margin = sec.bottom_margin = Cm(2.4)
    captioned = set()                      # a table caption becomes the table's own first row: Pages keeps a
    for tbl in d.element.body.findall(qn("w:tbl")):   # paragraph with the next paragraph but not with a table
        prev = tbl.getprevious()
        if prev is None or prev.tag != qn("w:p"):
            continue
        style = prev.find(".//" + qn("w:pStyle"))
        if style is None or style.get(qn("w:val")) != "TableCaption":
            continue
        t = next(x for x in d.tables if x._tbl is tbl)
        row = t.add_row()
        cell = row.cells[0].merge(row.cells[-1])
        cell._tc.remove(cell.paragraphs[0]._p)
        cell._tc.append(prev)                  # moves the caption paragraph, style and all
        borders = OxmlElement("w:tcBorders")
        for edge, val, sz in (("top", "single", "12"), ("bottom", "nil", "0")):
            e = OxmlElement(f"w:{edge}")
            for k, v in (("val", val), ("sz", sz), ("space", "0"), ("color", "111111")):
                e.set(qn(f"w:{k}"), v)
            borders.append(e)
        cell._tc.get_or_add_tcPr().append(borders)
        tbl.remove(row._tr)
        tbl.insert(list(tbl).index(tbl.findall(qn("w:tr"))[0]), row._tr)   # above the header row
        captioned.add(id(tbl))
    for t in d.tables:
        rule(t, header=1 if id(t._tbl) in captioned else 0)
        if id(t._tbl) in captioned:            # the caption row's top border is the table's opening rule
            t._tbl.tblPr.find(qn("w:tblBorders")).find(qn("w:top")).set(qn("w:val"), "nil")
        for row in t.rows:                  # a row never breaks across pages
            trpr = row._tr.get_or_add_trPr()
            cant = OxmlElement("w:cantSplit"); cant.set(qn("w:val"), "true"); trpr.append(cant)
            for cell in row.cells:
                for par in cell.paragraphs:
                    par.paragraph_format.keep_with_next = True
        for par in t.rows[-1].cells[0].paragraphs:   # but the table may still break after its last row
            par.paragraph_format.keep_with_next = False
    if "Source Code" in d.styles:          # pandoc adds this style itself, after the reference is read
        font(d.styles["Source Code"], MONO, 8.5, INK)
    for par in d.paragraphs:
        if par.style.name == "Source Code":        # a code block is set in mono and never breaks
            par.paragraph_format.keep_together = True
            for run in par.runs:
                run.font.name = MONO
                run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), MONO)
        if par.style.name in ("Table Caption", "Caption", "Image Caption"):
            par.paragraph_format.keep_with_next = True   # a caption stays with what follows it: the table, or the next text
        if par.style.name == "Captioned Figure" or par._element.findall(".//" + qn("w:drawing")):
            par.paragraph_format.keep_with_next = True   # an image stays with its caption
    text_width = d.sections[0].page_width - d.sections[0].left_margin - d.sections[0].right_margin
    for shape in d.inline_shapes:          # figures are authored at the text width: print them 1:1
        shape.height = int(shape.height * text_width / shape.width)
        shape.width = text_width
    d.core_properties.title = title if args else "Loomground Governance"
    d.save(out)
    from restyle_docx import restyle    # the Loomground palette, with no Word or pandoc colour left over
    restyle(out)
    print(out)


if __name__ == "__main__":
    import sys
    main(sys.argv[1:])
